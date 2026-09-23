"""
Text extraction utilities for resumes and job descriptions.
Supports .pdf, .docx, and .txt files (paths or file-like objects,
Supports .pdf, .docx, and .txt files (paths or file-like objects).
"""
from pathlib import Path

import pdfplumber
from docx import Document


def extract_text_from_pdf(file) -> str:
    """Extract text from a PDF (path string or file-like object)."""
    text_parts = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file) -> str:
    """Extract text from a .docx file (path string or file-like object)."""
    doc = Document(file)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Resumes often use tables for layout (skills grids, date columns, etc.)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text_from_txt(file) -> str:
    """Extract text from a plain text file (path string or file-like object)."""
    if hasattr(file, "read"):
        raw = file.read()
        return raw.decode("utf-8", errors="ignore") if isinstance(raw, bytes) else raw
    with open(file, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_text(file, filename: str = None) -> str:
    """
    Dispatch to the right extractor based on file extension.

    `file` can be a path (str/Path) or a file-like object.
    `filename` is only needed if `file` is a file-like object whose
    `.name` attribute doesn't carry the extension.
    """
    name = filename or getattr(file, "name", None) or str(file)
    ext = Path(name).suffix.lower()

    if hasattr(file, "seek"):
        file.seek(0)

    if ext == ".pdf":
        return extract_text_from_pdf(file)
    elif ext == ".docx":
        return extract_text_from_docx(file)
    elif ext == ".txt":
        return extract_text_from_txt(file)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Please use PDF, DOCX, or TXT.")
